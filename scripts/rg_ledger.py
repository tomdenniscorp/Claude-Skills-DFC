#!/usr/bin/env python3
"""
rg_ledger.py - Populate and verify the Rose Grange CBA loan schedule workbook
("Rose Grange - Loan Schedule - FY27 Statement Verification.xlsx").

The 'Statement Checks' tab is the single source of truth (one row per bank document).
Every run reads it into a data model, merges in newly parsed PDFs, runs the checks,
and then RE-RENDERS all three tabs (Statement Checks / Statement Processing /
Statement Month-End Schedule) so row references and formula ranges are always exact.
Manual notes (Statement Checks 'Notes' section and column P 'Comment') are preserved.

Commands
  ingest  <workbook> <pdf> [<pdf> ...] [--dry-run] [--force] [--out <path>]
          Parse PDFs, check them against the workbook, and (unless blocked) write.
          Prints a JSON report. Exit code 0 = written / nothing to do, 2 = blocked (escalate).
  rebuild <workbook> [--out <path>]
          Re-render the three tabs from the workbook's own Statement Checks data.
  verify  <workbook>
          Run AFTER recalculation (xlsx skill scripts/recalc.py). Scans for formula errors
          and independently recomputes Processing + Month-End in Python. Exit 0 = PASS.
"""
import argparse
import json
import shutil
import sys
from copy import deepcopy
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName

sys.path.insert(0, str(Path(__file__).resolve().parent))
from rg_parse import parse_pdf  # noqa: E402

SC, SP, ME = "Statement Checks", "Statement Processing", "Statement Month-End Schedule"
DEFAULT_LIMIT = 66_800_000.0
DEFAULT_LF_RATE = 0.014
TOL = 1.0
VERSION = "rg_ledger v1"

# ----------------------------------------------------------------------------- styles
ARIAL = "Arial"
F_TITLE = Font(name=ARIAL, size=14, bold=True, color="002060")
F_SUB = Font(name=ARIAL, size=11, bold=True)
F_SECTION = Font(name=ARIAL, size=11, bold=True, color="1F3864")
F_HEAD = Font(name=ARIAL, size=11, bold=True, color="FFFFFF")
F_BODY = Font(name=ARIAL, size=8)
F_INPUT = Font(name=ARIAL, size=8, color="0000FF")
F_GREY = Font(name=ARIAL, size=8, color="808080")
F_NOTE = Font(name=ARIAL, size=8, italic=True)
F_BOLD8 = Font(name=ARIAL, size=8, bold=True)
FILL_SECTION = PatternFill("solid", fgColor="D9E1F2")
FILL_HEAD = PatternFill("solid", fgColor="1F3864")
FILL_IMPLIED = PatternFill("solid", fgColor="FFF2CC")
FILL_LEGACY = PatternFill("solid", fgColor="FCE4D6")
FILL_TOTAL = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="BFBFBF")
B_TOP = Border(top=Side(style="thin", color="000000"))
FMT_DATE = "dd-mmm-yyyy"
FMT_CCY = '#,##0.00;(#,##0.00);"-"'
FMT_BAL = "#,##0"
FMT_RATE = "0.0000%"
GREEN = (PatternFill("solid", fgColor="C6EFCE"), Font(name=ARIAL, size=8, color="006100"))
RED = (PatternFill("solid", fgColor="FFC7CE"), Font(name=ARIAL, size=8, color="9C0006"))
AMBER = (PatternFill("solid", fgColor="FFEB9C"), Font(name=ARIAL, size=8, color="9C5700"))

GAP_TXT = "GAP - implied period inserted"
LF_GAP_TXT = "GAP - period not contiguous"
PENDING_TXT = "PENDING - no PDF on file"
LEGACY = "Legacy schedule - PDF pending"


def d2(x):
    if isinstance(x, datetime):
        return x.date()
    return x


def fmt(d):
    return d.strftime("%d-%b-%Y") if d else "?"


def q(sheet):
    return "'" + sheet + "'"


# ============================================================================= model I/O
def _literal(v):
    """Return literal text if v is a plain string or a ="text" formula, else None."""
    if v is None:
        return None
    if isinstance(v, str) and v.startswith('="') and v.endswith('"') and '&' not in v and "IF(" not in v:
        return v[2:-1]
    if isinstance(v, str) and not v.startswith("="):
        return v
    return None


def load_model(path):
    wb = openpyxl.load_workbook(path)
    ws = wb[SC]
    m = {"rollovers": [], "line_fees": [], "side": [], "other": [], "notes": [],
         "limit": DEFAULT_LIMIT, "line_fee_rate": DEFAULT_LF_RATE}
    if SP in wb.sheetnames:
        lim = wb[SP]["E3"].value
        if isinstance(lim, (int, float)) and lim > 0:
            m["limit"] = float(lim)
    section, in_data = None, False
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str):
            s = a.strip()
            key = None
            if s == "Rollover Notices":
                key = "rollovers"
            elif s == "Line Fee Invoices":
                key = "line_fees"
            elif s.startswith("$10m Side Tranche") or s.startswith("Side Tranches"):
                key = "side"
            elif s.startswith("Other Notices"):
                key = "other"
            elif s == "Summary":
                key = "summary"
            elif s.startswith("Notes"):
                key = "notes"
            elif s.startswith("Automated flags"):
                key = "auto"
            if key:
                section, in_data = key, False
                continue
            if s.startswith("Statement\nDate") or s.startswith("Statement Date") or s == "Document Date":
                in_data = True
                continue
        if section == "notes":
            if isinstance(a, str) and a.strip():
                m["notes"].append(a.strip().lstrip("•").strip())
            continue
        if section not in ("rollovers", "line_fees", "side", "other") or not in_data:
            continue
        if a is None:
            in_data = False
            continue
        row = [ws.cell(r, c).value for c in range(1, 17)]
        if section == "other":
            m["other"].append({"date": d2(row[0]), "description": row[1], "amount": row[2],
                               "notes": row[3], "source_doc": row[12]})
            continue
        rec = {"stmt_date": d2(row[0]), "start": d2(row[1]), "end": d2(row[2]),
               "balance": row[4], "rate": row[5], "stated": row[6],
               "prior_paid": row[10] if isinstance(row[10], (int, float)) else None,
               "source_doc": row[12], "debit_date": d2(row[13]) if isinstance(row[13], (date, datetime)) else None,
               "data_source": row[14] or "PDF", "comment": row[15],
               "continuity_override": None}
        lit = _literal(row[11])
        if lit and not lit.startswith("n/a - first"):
            rec["continuity_override"] = lit
        m[section].append(rec)
    for k in ("rollovers", "line_fees", "side"):
        m[k].sort(key=lambda x: (x["start"], x["end"]))
    for lf in m["line_fees"]:
        if lf["debit_date"] is None and lf["end"]:
            lf["debit_date"] = lf["end"] + timedelta(days=1)
    for ro in m["rollovers"]:
        ro["debit_date"] = ro["start"]
    return m


# ============================================================================= ingest
def _chk(report, name, status, detail):
    report["checks"].append({"check": name, "status": status, "detail": detail})
    if status == "FAIL":
        report["blocking"].append(f"{name}: {detail}")


def ingest(model, recs, force=False):
    """Merge parsed docs into model. Returns (new_model, report)."""
    m = deepcopy(model)
    rep = {"added": [], "replaced_legacy": [], "skipped": [], "checks": [], "blocking": [], "warnings": []}
    known_docs = {x.get("source_doc") for k in ("rollovers", "line_fees", "side", "other") for x in m[k]}

    for rec in sorted(recs, key=lambda r: (r.get("period_start") or r.get("effective_date") or r.get("doc_date") or date.min)):
        f = rec["file"]
        for c in rec["parse_checks"]:
            if c["status"] != "PASS":
                _chk(rep, f"[{f}] {c['check']}", c["status"], c["detail"])
                if c["status"] == "WARN":
                    rep["warnings"].append(f"{f}: {c['check']} - {c['detail']}")
        if rec["is_duplicate_copy"]:
            rep["skipped"].append(f"{f}: duplicate copy - ignored")
            continue
        if f in known_docs:
            rep["skipped"].append(f"{f}: already in workbook")
            continue
        fields_ok = all(rec.get(k) is not None for k in ("period_start", "period_end", "loan_balance", "rate", "stated_amount"))
        if not rec["parse_ok"] and not (force and fields_ok and rec["doc_type"] in ("rollover", "line_fee")):
            rep["skipped"].append(f"{f}: not ingested - parse checks failed (see blocking)")
            continue
        kind = rec["doc_type"]

        if kind == "rollover":
            new = {"stmt_date": rec["doc_date"], "start": rec["period_start"], "end": rec["period_end"],
                   "balance": rec["loan_balance"], "rate": rec["rate"], "stated": rec["stated_amount"],
                   "prior_paid": rec["prior_payment_confirmed"], "source_doc": f, "debit_date": rec["period_start"],
                   "data_source": "PDF", "comment": None, "continuity_override": None}
            same = [x for x in m["rollovers"] if x["start"] == new["start"]]
            if same:
                old = same[0]
                if str(old["data_source"]).startswith("Legacy"):
                    diffs = [k for k in ("end", "rate", "balance") if
                             (round(old[k], 6) if isinstance(old[k], float) else old[k]) !=
                             (round(new[k], 6) if isinstance(new[k], float) else new[k])]
                    if abs((old["stated"] or 0) - new["stated"]) > TOL:
                        diffs.append("stated")
                    _chk(rep, f"[{f}] PDF agrees to legacy-sourced row {fmt(old['start'])}",
                         "PASS" if not diffs else "WARN",
                         "match" if not diffs else "differs on " + ", ".join(diffs) + " - PDF figures used")
                    new["comment"] = old.get("comment")
                    m["rollovers"].remove(old)
                    rep["replaced_legacy"].append(f"{fmt(new['start'])}: legacy row replaced by {f}")
                else:
                    _chk(rep, f"[{f}] Conflicts with existing row for period starting {fmt(new['start'])}", "FAIL",
                         f"existing source {old['source_doc']}")
                    continue
            m["rollovers"].append(new)
            m["rollovers"].sort(key=lambda x: x["start"])
            i = m["rollovers"].index(new)
            if i > 0:
                prev = m["rollovers"][i - 1]
                if new["start"] == prev["end"]:
                    ok = abs(new["prior_paid"] - prev["stated"]) <= TOL
                    _chk(rep, f"[{f}] Continuity: interest paid {new['prior_paid']:,.2f} = prior period due "
                              f"{prev['stated']:,.2f}", "PASS" if ok else "FAIL",
                         "prior row: " + str(prev["source_doc"]))
                elif new["start"] > prev["end"]:
                    _chk(rep, f"[{f}] Continuity", "WARN",
                         f"GAP {fmt(prev['end'])} -> {fmt(new['start'])}: implied period will be inserted, "
                         f"rate back-solved from the {new['prior_paid']:,.2f} payment")
                    rep["warnings"].append(f"Gap {fmt(prev['end'])} to {fmt(new['start'])} - missing rollover notice")
                else:
                    _chk(rep, f"[{f}] Continuity", "FAIL",
                         f"overlaps prior period ending {fmt(prev['end'])}")
                if abs(new["rate"] - prev["rate"]) > 0.0075:
                    _chk(rep, f"[{f}] Rate movement vs prior period", "WARN",
                         f"{prev['rate']:.4%} -> {new['rate']:.4%}")
                if new["balance"] != prev["balance"]:
                    _chk(rep, f"[{f}] Facility balance vs prior period", "FAIL",
                         f"{prev['balance']:,.0f} -> {new['balance']:,.0f} - balance change needs manual treatment")
            if i < len(m["rollovers"]) - 1:
                nxt = m["rollovers"][i + 1]
                if nxt["start"] == new["end"] and nxt["prior_paid"] is not None:
                    ok = abs(nxt["prior_paid"] - new["stated"]) <= TOL
                    _chk(rep, f"[{f}] Confirmed by following notice", "PASS" if ok else "FAIL",
                         f"{nxt['prior_paid']:,.2f} vs {new['stated']:,.2f}")
            rep["added"].append(f"Rollover {fmt(new['start'])} - {fmt(new['end'])} @ {new['rate']:.4%} "
                                f"= {new['stated']:,.2f} ({f})")

        elif kind == "line_fee":
            new = {"stmt_date": rec["doc_date"], "start": rec["period_start"], "end": rec["period_end"],
                   "balance": rec["loan_balance"], "rate": rec["rate"], "stated": rec["stated_amount"],
                   "prior_paid": None, "source_doc": f, "debit_date": rec["effective_date"],
                   "data_source": "PDF", "comment": None, "continuity_override": None}
            same = [x for x in m["line_fees"] if x["start"] == new["start"]]
            if same:
                old = same[0]
                if str(old["data_source"]).startswith("Legacy"):
                    m["line_fees"].remove(old)
                    rep["replaced_legacy"].append(f"Line fee {fmt(new['start'])}: legacy row replaced by {f}")
                else:
                    _chk(rep, f"[{f}] Conflicts with existing line fee row starting {fmt(new['start'])}", "FAIL",
                         f"existing source {old['source_doc']}")
                    continue
            m["line_fees"].append(new)
            m["line_fees"].sort(key=lambda x: x["start"])
            i = m["line_fees"].index(new)
            if i > 0:
                prev = m["line_fees"][i - 1]
                ok = new["start"] == prev["end"] + timedelta(days=1)
                _chk(rep, f"[{f}] Line fee continuity (starts day after prior invoice)", "PASS" if ok else "WARN",
                     f"prior ended {fmt(prev['end'])}, this starts {fmt(new['start'])}")
            if abs(new["balance"] - m["limit"]) > 0.5:
                _chk(rep, f"[{f}] Line fee base = facility limit in workbook", "WARN",
                     f"{new['balance']:,.2f} vs {m['limit']:,.2f}")
            rep["added"].append(f"Line fee {fmt(new['start'])} - {fmt(new['end'])} = {new['stated']:,.2f}, "
                                f"debited {fmt(new['debit_date'])} ({f})")

        elif kind == "intent":
            m["other"].append({"date": rec["doc_date"], "description": rec.get("summary"),
                               "amount": rec.get("amount"),
                               "notes": f"Preliminary notice for the {fmt(rec.get('effective_date'))} rollover; "
                                        f"superseded by the final Rollover Notice.", "source_doc": f})
            rep["added"].append(f"Other notice (info only): {f}")
        elif kind == "break_cost":
            m["other"].append({"date": rec.get("effective_date") or rec["doc_date"], "description": rec.get("summary"),
                               "amount": rec.get("amount"),
                               "notes": "One-off cost, not prorated in the month-end schedule.", "source_doc": f})
            _chk(rep, f"[{f}] Break cost received", "WARN", f"AUD {rec.get('amount')} - journal separately")
            rep["added"].append(f"Other notice: break cost {rec.get('amount')} ({f})")
        else:  # repayment / drawdown / unknown -> human judgement on ledger treatment
            _chk(rep, f"[{f}] {kind} document", "FAIL",
                 f"{rec.get('summary') or 'unrecognised document'} - changes the facility; needs manual ledger "
                 f"treatment (balance-change split or side-tranche section). Re-run with --force to record it under "
                 f"Other Notices only.")
            m["other"].append({"date": rec.get("effective_date") or rec["doc_date"], "description": rec.get("summary"),
                               "amount": rec.get("amount"), "notes": "Added by --force run; ledger impact reviewed manually.",
                               "source_doc": f})
    m["other"].sort(key=lambda x: (x["date"] or date.min))
    rep["status"] = "ESCALATE" if rep["blocking"] else ("NO CHANGE" if not (rep["added"] or rep["replaced_legacy"]) else "ROUTINE")
    return m, rep


# ============================================================================= rendering helpers
def _hdr(ws, row, labels, height=30):
    for i, lab in enumerate(labels, 1):
        c = ws.cell(row, i, lab)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
    ws.row_dimensions[row].height = height


def _section(ws, row, text, ncols):
    for i in range(1, ncols + 1):
        ws.cell(row, i).fill = FILL_SECTION
    c = ws.cell(row, 1, text)
    c.font = F_SECTION


def _put(ws, row, col, val, fmt_=None, font=F_BODY, fill=None):
    c = ws.cell(row, col, val)
    c.font = font
    if fmt_:
        c.number_format = fmt_
    if fill:
        c.fill = fill
    return c


def _status_cf(ws, rng):
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"PASS"'], fill=GREEN[0], font=GREEN[1]))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"FAIL"'], fill=RED[0], font=RED[1]))
    first = rng.split(":")[0].replace("$", "")
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR(LEFT({first},3)="GAP",LEFT({first},7)="PENDING",'
                                                            f'LEFT({first},7)="Partial",LEFT({first},5)="CHECK")'],
                                                   fill=AMBER[0], font=AMBER[1]))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Full month"'], fill=GREEN[0], font=GREEN[1]))


def _missing_line_fee_quarters(lfs):
    out = []
    for a, b in zip(lfs, lfs[1:]):
        if b["start"] != a["end"] + timedelta(days=1):
            out.append((a["end"] + timedelta(days=1), b["start"] - timedelta(days=1)))
    return out


# ============================================================================= render: Statement Checks
SC_COLS = ["Statement\nDate", "Period\nStart", "Period\nEnd", "Days", "Loan\nBalance", "Rate",
           "Stated Amount\n(AUD)", "Recalculated\nAmount (AUD)", "Variance\n($)", "Calc\nCheck",
           "Prior-Period\nPayment Confirmed", "Continuity\nCheck", "Source Document", "Debit\nDate",
           "Data\nSource", "Comment\n(manual)"]
NC = len(SC_COLS)


def render_checks(ws, m, today):
    lay = {"ro": {}, "lf": {}, "side": {}}
    ro, lf, sd = m["rollovers"], m["line_fees"], m["side"]
    first = min([x["start"] for x in ro + lf + sd] or [today])
    last = max([x["end"] for x in ro + lf + sd] or [today])
    ws["A1"] = "ROSE GRANGE COMMERCIAL TRUST (052) – CBA FACILITY B"
    ws["A1"].font = F_TITLE
    ws["A2"] = f"Statement Checks – document-by-document extraction & verification ({first:%b-%Y} to {last:%b-%Y})"
    ws["A2"].font = F_SUB
    ws["A3"] = ("Deal Ref: ROSE GRANGE COM BILAT AUD66.8M 15DEC17 | Bank Ref: AUSNA1000072 | Rate Basis: Actual/365 | "
                "Blue = extracted from bank documents; black = formula. Column P is free-text and preserved each run.")
    ws["A3"].font = F_BODY
    ws["A4"] = f"Last updated {today:%d-%b-%Y} by {VERSION} (rose-grange-statement-parse skill)"
    ws["A4"].font = F_GREY
    r = 6

    def block(title, rows, kind):
        nonlocal r
        _section(ws, r, title, NC)
        r += 1
        _hdr(ws, r, SC_COLS)
        r += 1
        r0 = r
        for i, x in enumerate(rows):
            legacy = str(x["data_source"]).startswith("Legacy")
            fill = FILL_LEGACY if legacy else None
            _put(ws, r, 1, x["stmt_date"], FMT_DATE, F_INPUT, fill)
            _put(ws, r, 2, x["start"], FMT_DATE, F_INPUT, fill)
            _put(ws, r, 3, x["end"], FMT_DATE, F_INPUT, fill)
            _put(ws, r, 4, f"=C{r}-B{r}+1" if kind == "lf" else f"=C{r}-B{r}", "0", F_BODY, fill)
            _put(ws, r, 5, x["balance"], FMT_BAL, F_INPUT, fill)
            _put(ws, r, 6, x["rate"], FMT_RATE, F_INPUT, fill)
            _put(ws, r, 7, x["stated"], FMT_CCY, F_INPUT, fill)
            _put(ws, r, 8, f"=E{r}*F{r}*D{r}/365", FMT_CCY, F_BODY, fill)
            _put(ws, r, 9, f"=G{r}-H{r}", FMT_CCY, F_BODY, fill)
            _put(ws, r, 10, f'=IF(ABS(I{r})<={TOL},"PASS","FAIL")', None, F_BODY, fill)
            _put(ws, r, 11, x["prior_paid"], FMT_CCY, F_INPUT, fill)
            if x["continuity_override"]:
                lval = f'="{x["continuity_override"]}"'
            elif kind == "ro":
                if i == 0:
                    lval = '="n/a - first statement in this set"'
                else:
                    lval = (f'=IF(B{r}<>C{r-1},"{GAP_TXT}",IF(ISNUMBER(K{r}),IF(ABS(K{r}-G{r-1})<={TOL},"PASS","FAIL"),'
                            f'"{PENDING_TXT}"))')
            elif kind == "lf":
                lval = '="n/a - first fee invoice in this set"' if i == 0 else \
                    f'=IF(B{r}=C{r-1}+1,"PASS","{LF_GAP_TXT}")'
            else:
                lval = '="n/a - separate tranche"'
            _put(ws, r, 12, lval, None, F_BODY, fill)
            _put(ws, r, 13, x["source_doc"], None, F_GREY, fill)
            _put(ws, r, 14, x.get("debit_date"), FMT_DATE, F_INPUT, fill)
            _put(ws, r, 15, x["data_source"], None, F_GREY, fill)
            _put(ws, r, 16, x.get("comment"), None, F_BODY, fill)
            r += 1
        r1 = r - 1
        if rows:
            _status_cf(ws, f"J{r0}:J{r1}")
            _status_cf(ws, f"L{r0}:L{r1}")
        r += 1
        return r0, r1

    lay["ro"] = block("Rollover Notices", ro, "ro")
    lay["lf"] = block("Line Fee Invoices", lf, "lf")
    lay["side"] = block("Side Tranches (separate concurrent drawdowns – excluded from the main rollover chain)", sd, "side")

    _section(ws, r, "Other Notices (repayments, break costs, intent notices) – recorded for the record, not part of the "
                    "recalculation chain", NC)
    r += 1
    _hdr(ws, r, ["Document Date", "Description", "Amount (AUD)", "Notes"] + [""] * 8 + ["Source Document"], 20)
    r += 1
    for o in m["other"]:
        _put(ws, r, 1, o["date"], FMT_DATE, F_INPUT)
        _put(ws, r, 2, o["description"])
        _put(ws, r, 3, o["amount"], FMT_CCY, F_INPUT)
        _put(ws, r, 4, o["notes"])
        _put(ws, r, 13, o["source_doc"], None, F_GREY)
        r += 1
    r += 1

    _section(ws, r, "Summary", NC)
    r += 1
    (a0, a1), (b0, b1) = lay["ro"], lay["lf"]
    lines = []
    if a1 >= a0:
        lines += [f'="Rollover calc checks: "&COUNTIF(J{a0}:J{a1},"PASS")&" of "&COUNTA(J{a0}:J{a1})&" PASS (tolerance ${TOL:.0f})"',
                  f'="Rollover continuity: "&COUNTIF(L{a0}:L{a1},"PASS")&" PASS, "&COUNTIF(L{a0}:L{a1},"FAIL")&" FAIL, "'
                  f'&COUNTIF(L{a0}:L{a1},"{GAP_TXT}")&" GAP (implied period inserted), "&COUNTIF(L{a0}:L{a1},"{PENDING_TXT}")'
                  f'&" PENDING (legacy figures, PDF not on file)"']
    if b1 >= b0:
        lines += [f'="Line fee calc checks: "&COUNTIF(J{b0}:J{b1},"PASS")&" of "&COUNTA(J{b0}:J{b1})&" PASS; continuity: "'
                  f'&COUNTIF(L{b0}:L{b1},"{LF_GAP_TXT}")&" gap(s)"']
    for ln in lines:
        _put(ws, r, 1, ln, None, F_BOLD8)
        r += 1
    r += 1

    _section(ws, r, "Notes (manual – preserved each run; edit freely)", NC)
    r += 1
    for n in m["notes"]:
        _put(ws, r, 1, "• " + n)
        r += 1
    r += 1

    _section(ws, r, "Automated flags (regenerated each run – do not edit)", NC)
    r += 1
    for n in auto_flags(m, today):
        _put(ws, r, 1, "• " + n, None, F_NOTE)
        r += 1

    widths = [12, 12, 12, 7, 13, 10, 15, 15, 10, 9, 16, 30, 52, 12, 26, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "B6"
    return lay


def auto_flags(m, today):
    out = []
    ro, lf = m["rollovers"], m["line_fees"]
    for a, b in zip(ro, ro[1:]):
        if b["start"] > a["end"]:
            out.append(f"{fmt(a['end'])} to {fmt(b['start'])}: no rate-set notice on file – implied period inserted in "
                       f"'Statement Processing', interest taken from the {fmt(b['stmt_date'])} notice's payment confirmation.")
    leg = [x for x in ro + lf if str(x["data_source"]).startswith("Legacy")]
    for x in leg:
        out.append(f"{fmt(x['start'])} to {fmt(x['end'])} @ {x['rate']:.4%}: figures taken from the legacy Loan Schedule – "
                   f"replace by ingesting the CBA PDF when available.")
    if ro:
        z = ro[-1]
        out.append(f"Latest period {fmt(z['start'])} to {fmt(z['end'])} (AUD {z['stated']:,.2f}) is projected – it is "
                   f"confirmed by the next rollover notice, expected on/after {fmt(z['end'])}.")
    for s, e in _missing_line_fee_quarters(lf):
        out.append(f"Line fee quarter {fmt(s)} to {fmt(e)}: no invoice on file (accrual still recognised in the "
                   f"Month-End schedule; cash column will show nil).")
    if lf:
        z = lf[-1]
        out.append(f"Next line fee invoice expected for {fmt(z['end'] + timedelta(days=1))} onwards "
                   f"(quarter ending ~22nd, issued mid-month).")
    return out


# ============================================================================= render: Statement Processing
SP_COLS = ["Bill Start\nDate", "Bill End\nDate", "Days", "Interest\nRate", "Loan\nBalance",
           "Interest – stated\n(AUD)", "Interest – realised\n(AUD, used for accrual)", "Interest Paid at start\n(prior period, AUD)",
           "Line fee\nrate", "Line Fee on limit\n(bill period, AUD)", "Source Document", "Status"]


def render_processing(ws, m, lay):
    sc = q(SC)
    ws["A1"] = "ROSE GRANGE COMMERCIAL TRUST (052) – CBA FACILITY B"
    ws["A1"].font = F_TITLE
    ws["A2"] = "Statement Processing – continuous bill-period ledger, formula-linked to 'Statement Checks' (rebuilt each run)"
    ws["A2"].font = F_SUB
    ws["A3"] = "Facility limit (line fee base):"
    ws["A3"].font = F_BOLD8
    _put(ws, 3, 5, m["limit"], FMT_BAL, F_INPUT)
    _hdr(ws, 4, SP_COLS, 42)
    r = 5
    ro = m["rollovers"]
    a0 = lay["ro"][0]
    main = []
    for i, x in enumerate(ro):
        c = a0 + i
        if i > 0 and x["start"] > ro[i - 1]["end"]:
            cp = c - 1
            confirmed = x["prior_paid"] is not None
            _put(ws, r, 1, f"={sc}!C{cp}", FMT_DATE, fill=FILL_IMPLIED)
            _put(ws, r, 2, f"={sc}!B{c}", FMT_DATE, fill=FILL_IMPLIED)
            _put(ws, r, 3, f"=B{r}-A{r}", "0", fill=FILL_IMPLIED)
            _put(ws, r, 4, f"=IF(E{r}*C{r}=0,0,F{r}*365/(E{r}*C{r}))", FMT_RATE, fill=FILL_IMPLIED)
            _put(ws, r, 5, f"={sc}!E{c}", FMT_BAL, fill=FILL_IMPLIED)
            _put(ws, r, 6, f"=IF(ISNUMBER({sc}!K{c}),{sc}!K{c},0)", FMT_CCY, fill=FILL_IMPLIED)
            _put(ws, r, 7, f"=F{r}", FMT_CCY, fill=FILL_IMPLIED)
            _put(ws, r, 8, f"={sc}!G{cp}", FMT_CCY, fill=FILL_IMPLIED)
            _put(ws, r, 9, m["line_fee_rate"], FMT_RATE, F_INPUT, FILL_IMPLIED)
            _put(ws, r, 10, f"=$E$3*I{r}*C{r}/365", FMT_CCY, fill=FILL_IMPLIED)
            _put(ws, r, 11, f"Implied – no rate-set notice; interest confirmed by {fmt(x['stmt_date'])} notice", None,
                 F_GREY, FILL_IMPLIED)
            _put(ws, r, 12, "Implied period – rate back-solved from confirmed payment" if confirmed else
                 "UNCONFIRMED GAP – no notice and no payment confirmation", None, F_BODY, FILL_IMPLIED)
            main.append(r)
            r += 1
        legacy = str(x["data_source"]).startswith("Legacy")
        fill = FILL_LEGACY if legacy else None
        nxt = ro[i + 1] if i + 1 < len(ro) else None
        _put(ws, r, 1, f"={sc}!B{c}", FMT_DATE, fill=fill)
        _put(ws, r, 2, f"={sc}!C{c}", FMT_DATE, fill=fill)
        _put(ws, r, 3, f"=B{r}-A{r}", "0", fill=fill)
        _put(ws, r, 4, f"={sc}!F{c}", FMT_RATE, fill=fill)
        _put(ws, r, 5, f"={sc}!E{c}", FMT_BAL, fill=fill)
        _put(ws, r, 6, f"={sc}!G{c}", FMT_CCY, fill=fill)
        if nxt and nxt["start"] == x["end"]:
            _put(ws, r, 7, f"=IF(ISNUMBER({sc}!K{c+1}),{sc}!K{c+1},F{r})", FMT_CCY, fill=fill)
        else:
            _put(ws, r, 7, f"=F{r}", FMT_CCY, fill=fill)
        _put(ws, r, 8, f'=IF(ISNUMBER({sc}!K{c}),{sc}!K{c},0)', FMT_CCY, fill=fill)
        _put(ws, r, 9, m["line_fee_rate"], FMT_RATE, F_INPUT, fill)
        _put(ws, r, 10, f"=$E$3*I{r}*C{r}/365", FMT_CCY, fill=fill)
        _put(ws, r, 11, f"={sc}!M{c}", None, F_GREY, fill)
        if nxt is None:
            st = "Projected – awaiting next rollover notice"
        elif nxt["start"] != x["end"]:
            st = "Rate-set notice only – following period implied"
        elif nxt["prior_paid"] is None:
            st = "Not yet confirmed – following row is legacy-sourced"
        elif abs(nxt["prior_paid"] - x["stated"]) <= TOL:
            st = "Confirmed (rate-set + next notice payment match)"
        else:
            st = "Realised ≠ stated (mid-period balance change) – realised amount used"
        if legacy:
            st = "Legacy schedule figures – " + st[0].lower() + st[1:]
        _put(ws, r, 12, st, None, F_BODY, fill)
        main.append(r)
        r += 1
    m0, m1 = (main[0], main[-1]) if main else (5, 4)
    r += 1
    side_rows = []
    if m["side"]:
        _put(ws, r, 1, "Side tranches (separate concurrent drawdowns – included in interest accrual, excluded from "
                       "Days Captured / balance lookups)", None, F_SECTION)
        r += 1
        s0c = lay["side"][0]
        for j, x in enumerate(m["side"]):
            c = s0c + j
            _put(ws, r, 1, f"={sc}!B{c}", FMT_DATE)
            _put(ws, r, 2, f"={sc}!C{c}", FMT_DATE)
            _put(ws, r, 3, f"=B{r}-A{r}", "0")
            _put(ws, r, 4, f"={sc}!F{c}", FMT_RATE)
            _put(ws, r, 5, f"={sc}!E{c}", FMT_BAL)
            _put(ws, r, 6, f"={sc}!G{c}", FMT_CCY)
            _put(ws, r, 7, f"=F{r}", FMT_CCY)
            _put(ws, r, 8, 0, FMT_CCY)
            _put(ws, r, 9, 0, FMT_RATE, F_INPUT)
            _put(ws, r, 10, 0, FMT_CCY)
            _put(ws, r, 11, f"={sc}!M{c}", None, F_GREY)
            _put(ws, r, 12, "Separate tranche – no line fee", None)
            side_rows.append(r)
            r += 1
    r += 1
    _put(ws, r, 1, "Notes: 'Interest – realised' = the amount the NEXT notice confirms was paid (falls back to the stated "
                   "amount if not yet confirmed). It differs from 'stated' only where the balance changed mid-period. "
                   "Amber rows = implied periods; orange rows = legacy-sourced figures awaiting the CBA PDF. Break costs "
                   "and principal movements are recorded in 'Statement Checks' (Other Notices), not here.", None, F_NOTE)
    widths = [12, 12, 7, 10, 13, 15, 17, 17, 9, 15, 55, 48]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A5"
    return {"main": (m0, m1), "side": (side_rows[0], side_rows[-1]) if side_rows else None}


# ============================================================================= render: Month-End
ME_COLS = ["Month\nEnd", "Days", "Loan Balance\n(at month end)", "Interest Rate\n(at month end)", "Interest\n(accrual, AUD)",
           "Interest Paid\n(cash, AUD)", "Line fee\nrate", "Line Fee\n(accrual, AUD)", "Line Fee Paid\n(cash, AUD)",
           "Total\nAccrual (AUD)", "Days\nCaptured", "Coverage", "FY"]


def _eom(d):
    nxt = date(d.year + (d.month == 12), d.month % 12 + 1, 1)
    return nxt - timedelta(days=1)


def render_month_end(ws, m, lay, play):
    sp, sc = q(SP), q(SC)
    m0, m1 = play["main"]
    ranges = [(m0, m1)] + ([play["side"]] if play["side"] else [])
    ro = m["rollovers"]
    first_month = _eom(ro[0]["start"]) if ro else _eom(date.today())
    last_end = max([x["end"] for x in ro + m["side"]] or [first_month])
    last_month = _eom(last_end - timedelta(days=1))
    ws["A1"] = "ROSE GRANGE COMMERCIAL TRUST (052) – CBA FACILITY B"
    ws["A1"].font = F_TITLE
    ws["A2"] = "Statement Month-End Schedule – calendar-month accrual / cash split for accounting journals"
    ws["A2"].font = F_SUB
    ws["A3"] = (f"Sourced from 'Statement Processing' – covers {first_month:%b-%Y} to {last_month:%b-%Y}. Interest accrual "
                f"prorates each bill period's realised interest by days in month; line fee accrues daily on the facility limit.")
    ws["A3"].font = F_BODY
    _hdr(ws, 5, ME_COLS, 42)
    r = 6
    ms, me = "(EOMONTH($A{r},-1)+1)", "($A{r}+1)"

    def rng(col, a, b):
        return f"{sp}!${col}${a}:${col}${b}"

    def overlap(a, b, rr):
        S, E = rng("A", a, b), rng("B", a, b)
        s_, e_ = ms.format(r=rr), me.format(r=rr)
        inside = f"({E}>{s_})*({S}<{e_})"
        ov = f"(({E}<{e_})*{E}+({E}>={e_})*{e_}-(({S}>{s_})*{S}+({S}<={s_})*{s_}))"
        return inside, ov

    mo = first_month
    while mo <= last_month:
        _put(ws, r, 1, mo, FMT_DATE)
        _put(ws, r, 2, f"=DAY($A{r})", "0")
        _put(ws, r, 3, f'=IFERROR(INDEX({rng("E", m0, m1)},MATCH($A{r},{rng("A", m0, m1)},1)),"n/a")', FMT_BAL)
        _put(ws, r, 4, f'=IFERROR(INDEX({rng("D", m0, m1)},MATCH($A{r},{rng("A", m0, m1)},1)),"n/a")', FMT_RATE)
        parts_e, parts_f = [], []
        for a, b in ranges:
            inside, ov = overlap(a, b, r)
            parts_e.append(f"SUMPRODUCT({inside}*{ov}*{rng('G', a, b)}/{rng('C', a, b)})")
            E = rng("B", a, b)
            parts_f.append(f"SUMPRODUCT(({E}>={ms.format(r=r)})*({E}<{me.format(r=r)})*{rng('G', a, b)})")
        _put(ws, r, 5, "=" + "+".join(parts_e), FMT_CCY)
        _put(ws, r, 6, "=" + "+".join(parts_f), FMT_CCY)
        _put(ws, r, 7, m["line_fee_rate"], FMT_RATE, F_INPUT)
        _put(ws, r, 8, f"={sp}!$E$3*G{r}*B{r}/365", FMT_CCY)
        b0, b1 = lay["lf"]
        if b1 >= b0:
            N, G = f"{sc}!$N${b0}:$N${b1}", f"{sc}!$G${b0}:$G${b1}"
            _put(ws, r, 9, f"=SUMPRODUCT(({N}>={ms.format(r=r)})*({N}<{me.format(r=r)})*{G})", FMT_CCY)
        else:
            _put(ws, r, 9, 0, FMT_CCY)
        _put(ws, r, 10, f"=E{r}+H{r}", FMT_CCY)
        inside, ov = overlap(m0, m1, r)
        _put(ws, r, 11, f"=SUMPRODUCT({inside}*{ov})", "0")
        _put(ws, r, 12, f'=IF(K{r}=B{r},"Full month",IF(K{r}>B{r},"CHECK - overlapping periods ("&K{r}&" days)",'
                        f'"Partial - "&K{r}&" of "&B{r}&" days (missing / future statement)"))')
        _put(ws, r, 13, f'="FY"&RIGHT(YEAR($A{r})+IF(MONTH($A{r})>6,1,0),2)')
        r += 1
        mo = _eom(mo + timedelta(days=1))
    d0, d1 = 6, r - 1
    _status_cf(ws, f"L{d0}:L{d1}")
    for col in (1, 5, 6, 8, 9, 10):
        c = ws.cell(r, col)
        c.border, c.fill = B_TOP, FILL_TOTAL
    _put(ws, r, 1, "Total", None, F_BOLD8, FILL_TOTAL)
    for col in "EFHIJ":
        ci = openpyxl.utils.column_index_from_string(col)
        _put(ws, r, ci, f"=SUM({col}{d0}:{col}{d1})", FMT_CCY, F_BOLD8, FILL_TOTAL)
    tot = r
    r += 2
    _put(ws, r, 1, "Reconciliation checks", None, F_SUB)
    r += 1
    g_terms = "+".join(f"SUM({rng('G', a, b)})" for a, b in ranges)
    _put(ws, r, 1, "Interest accrual (Σ months) less realised interest (Σ ledger periods) – should be nil")
    _put(ws, r, 10, f"=ROUND(E{tot}-({g_terms}),2)", FMT_CCY)
    _put(ws, r, 12, f'=IF(ABS(J{r})<0.01,"PASS","FAIL")')
    rc1 = r
    r += 1
    _put(ws, r, 1, "Interest cash (Σ months) less realised interest (Σ ledger periods) – should be nil")
    _put(ws, r, 10, f"=ROUND(F{tot}-({g_terms}),2)", FMT_CCY)
    _put(ws, r, 12, f'=IF(ABS(J{r})<0.01,"PASS","FAIL")')
    r += 1
    if lay["lf"][1] >= lay["lf"][0]:
        b0, b1 = lay["lf"]
        _put(ws, r, 1, "Line fee cash (Σ months) less line fee invoices on file – should be nil")
        _put(ws, r, 10, f"=ROUND(I{tot}-SUMPRODUCT(({sc}!$N${b0}:$N${b1}<=$A${d1})*{sc}!$G${b0}:$G${b1}),2)", FMT_CCY)
        _put(ws, r, 12, f'=IF(ABS(J{r})<0.01,"PASS","FAIL")')
        r += 1
    _status_cf(ws, f"L{rc1}:L{r-1}")
    r += 1
    _put(ws, r, 1, "Interest Paid / Line Fee Paid = amounts debited by the bank in that calendar month (interest on each "
                   "bill end date; line fee on the invoice debit date). Loan Balance / Rate = the bill period in force at "
                   "month end. Months flagged Partial are not fully evidenced by rollover notices yet.", None, F_NOTE)
    widths = [12, 6, 14, 12, 15, 15, 9, 14, 15, 15, 9, 40, 7]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "B6"
    return {"data": (d0, d1), "total": tot}


# ============================================================================= write
def render(path_in, m, path_out, today=None):
    today = today or date.today()
    wb = openpyxl.load_workbook(path_in)
    order = [SC, SP, ME]
    for name in order:
        if name in wb.sheetnames:
            del wb[name]
    for name in list(wb.sheetnames):
        if name.startswith("__placeholder"):
            del wb[name]
    ws_sc = wb.create_sheet(SC, 0)
    ws_sp = wb.create_sheet(SP, 1)
    ws_me = wb.create_sheet(ME, 2)
    lay = render_checks(ws_sc, m, today)
    play = render_processing(ws_sp, m, lay)
    mlay = render_month_end(ws_me, m, lay, play)
    for nm in ("RG_MAIN", "RG_SIDE", "RG_MONTHS"):
        if nm in wb.defined_names:
            del wb.defined_names[nm]
    m0, m1 = play["main"]
    wb.defined_names["RG_MAIN"] = DefinedName("RG_MAIN", attr_text=f"{q(SP)}!$A${m0}:$L${m1}")
    if play["side"]:
        wb.defined_names["RG_SIDE"] = DefinedName("RG_SIDE", attr_text=f"{q(SP)}!$A${play['side'][0]}:$L${play['side'][1]}")
    d0, d1 = mlay["data"]
    wb.defined_names["RG_MONTHS"] = DefinedName("RG_MONTHS", attr_text=f"{q(ME)}!$A${d0}:$M${d1}")
    for ws in (ws_sc, ws_sp, ws_me):
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_view.showGridLines = False
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == SC
    wb.save(path_out)
    return {"checks": lay, "processing": play, "month_end": mlay}


def backup(path):
    p = Path(path)
    bdir = p.parent / "_backups"
    bdir.mkdir(exist_ok=True)
    dst = bdir / f"{p.stem} - backup {datetime.now():%Y%m%d-%H%M%S}{p.suffix}"
    shutil.copy2(p, dst)
    return str(dst)


# ============================================================================= verify
ERR = ("#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "Err:")


def _range(wb, name):
    dn = wb.defined_names.get(name)
    if not dn:
        return None
    ref = dn.attr_text.split("!")[1].replace("$", "")
    a, b = ref.split(":")
    return int("".join(ch for ch in a if ch.isdigit())), int("".join(ch for ch in b if ch.isdigit()))


def verify(path):
    wbv = openpyxl.load_workbook(path, data_only=True)
    res = {"errors": [], "checks": []}

    def add(name, ok, detail=""):
        res["checks"].append({"check": name, "status": "PASS" if ok else "FAIL", "detail": detail})

    add("Exactly the three ledger tabs present first", wbv.sheetnames[:3] == [SC, SP, ME], str(wbv.sheetnames))
    uncached = 0
    wbf = openpyxl.load_workbook(path)
    for ws in wbv.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(ERR):
                    res["errors"].append(f"{ws.title}!{c.coordinate} {c.value}")
                if c.value is None and isinstance(wbf[ws.title][c.coordinate].value, str) and \
                        wbf[ws.title][c.coordinate].value.startswith("="):
                    uncached += 1
    add("No formula errors", not res["errors"], "; ".join(res["errors"][:10]))
    add("Formulas have calculated values (recalc has run)", uncached < 5, f"{uncached} formula cells without values")

    sp, me, sc = wbv[SP], wbv[ME], wbv[SC]
    periods = []
    m0, m1 = _range(wbv, "RG_MAIN")
    main = []
    for r in range(m0, m1 + 1):
        s, e, dd, g = (d2(sp.cell(r, 1).value), d2(sp.cell(r, 2).value), sp.cell(r, 3).value, sp.cell(r, 7).value)
        main.append((s, e, dd, g))
    periods += main
    sr = _range(wbv, "RG_SIDE")
    if sr:
        for r in range(sr[0], sr[1] + 1):
            periods.append((d2(sp.cell(r, 1).value), d2(sp.cell(r, 2).value), sp.cell(r, 3).value, sp.cell(r, 7).value))
    cont = all(main[i + 1][0] == main[i][1] for i in range(len(main) - 1))
    add("Processing main ledger is continuous (each start = prior end)", cont)
    add("Processing days = end - start", all((e - s).days == dd for s, e, dd, _ in periods))

    lfs = []
    for r in range(1, sc.max_row + 1):
        if sc.cell(r, 1).value == "Line Fee Invoices":
            rr = r + 2
            while isinstance(sc.cell(rr, 1).value, (date, datetime)):
                lfs.append((d2(sc.cell(rr, 14).value), sc.cell(rr, 7).value))
                rr += 1
            break
    limit = sp["E3"].value
    d0, d1 = _range(wbv, "RG_MONTHS")
    bad = []
    for r in range(d0, d1 + 1):
        mo = d2(me.cell(r, 1).value)
        ms, mend = date(mo.year, mo.month, 1), mo + timedelta(days=1)
        acc = cash = 0.0
        days_cap = 0
        for idx, (s, e, dd, g) in enumerate(periods):
            ov = max(0, (min(e, mend) - max(s, ms)).days)
            if ov and dd:
                acc += g * ov / dd
            if ms <= e < mend:
                cash += g
            if idx < len(main):
                days_cap += ov
        lf_acc = limit * (me.cell(r, 7).value or 0) * mo.day / 365
        lf_cash = sum(a for dt, a in lfs if dt and ms <= dt < mend)
        exp = {5: acc, 6: cash, 8: lf_acc, 9: lf_cash, 11: days_cap}
        for col, v in exp.items():
            got = me.cell(r, col).value
            if not isinstance(got, (int, float)) or abs(got - v) > 0.01:
                bad.append(f"{mo:%b-%Y} col {openpyxl.utils.get_column_letter(col)}: sheet {got} vs python {v:.2f}")
    add("Month-End recomputed independently in Python (accrual, cash, line fee, days)", not bad, "; ".join(bad[:8]))
    recon = [me.cell(r, 12).value for r in range(d1 + 2, me.max_row + 1) if me.cell(r, 12).value in ("PASS", "FAIL")]
    add("Month-End reconciliation rows all PASS", recon and all(v == "PASS" for v in recon), str(recon))
    jfails = [f"{SC}!J{r}" for r in range(1, sc.max_row + 1) if sc.cell(r, 10).value == "FAIL"]
    lfails = [f"{SC}!L{r}" for r in range(1, sc.max_row + 1) if sc.cell(r, 12).value == "FAIL"]
    add("Statement Checks: no Calc Check FAIL", not jfails, ", ".join(jfails))
    add("Statement Checks: no Continuity FAIL", not lfails, ", ".join(lfails))
    partial = [f"{d2(me.cell(r, 1).value):%b-%Y}: {me.cell(r, 12).value}" for r in range(d0, d1 + 1)
               if me.cell(r, 12).value != "Full month"]
    res["partial_months"] = partial
    res["status"] = "PASS" if all(c["status"] == "PASS" for c in res["checks"]) else "FAIL"
    return res


# ============================================================================= CLI
def _jd(o):
    if isinstance(o, (date, datetime)):
        return o.isoformat()
    return str(o)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("ingest")
    a.add_argument("workbook")
    a.add_argument("pdfs", nargs="+")
    a.add_argument("--dry-run", action="store_true")
    a.add_argument("--force", action="store_true")
    a.add_argument("--out")
    b = sub.add_parser("rebuild")
    b.add_argument("workbook")
    b.add_argument("--out")
    c = sub.add_parser("verify")
    c.add_argument("workbook")
    args = ap.parse_args()

    if args.cmd == "verify":
        res = verify(args.workbook)
        print(json.dumps(res, indent=2, default=_jd))
        sys.exit(0 if res["status"] == "PASS" else 1)

    model = load_model(args.workbook)
    if args.cmd == "rebuild":
        out = args.out or args.workbook
        if out == args.workbook:
            print("backup:", backup(args.workbook))
        render(args.workbook, model, out)
        print(json.dumps({"status": "REBUILT", "out": out, "rollovers": len(model["rollovers"]),
                          "line_fees": len(model["line_fees"])}, indent=2))
        return

    recs = [parse_pdf(p) for p in args.pdfs]
    new_model, rep = ingest(model, recs, force=args.force)
    rep["parsed"] = [{k: v for k, v in r.items() if k not in ("parse_checks",)} for r in recs]
    blocked = bool(rep["blocking"]) and not args.force
    if args.dry_run or blocked or rep["status"] == "NO CHANGE":
        rep["written"] = False
        if blocked:
            rep["note"] = "Blocked - nothing written. Resolve the blocking items or re-run with --force."
    else:
        out = args.out or args.workbook
        if out == args.workbook:
            rep["backup"] = backup(args.workbook)
        render(args.workbook, new_model, out)
        rep["written"] = True
        rep["out"] = out
    print(json.dumps(rep, indent=2, default=_jd))
    sys.exit(2 if blocked else 0)


if __name__ == "__main__":
    main()
